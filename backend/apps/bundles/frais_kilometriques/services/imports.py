from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook

from apps.bundles.frais_kilometriques.models import JourTravaille, SimulationFraisKilometriques
from apps.bundles.frais_kilometriques.services.distances import get_or_compute_distance
from apps.bundles.frais_kilometriques.services.openrouteservice import OpenRouteServiceError


class ImportFraisKilometriquesError(Exception):
    pass


def _normalize_text(value: str) -> str:
    return value.strip().lower()


def _parse_excel_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        raw = value.strip()
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(raw, fmt).date()
            except ValueError:
                continue
    raise ImportFraisKilometriquesError("Format de date invalide dans le fichier Excel.")


def import_calendar_xlsx(*, simulation: SimulationFraisKilometriques, file_bytes: bytes) -> dict:
    workbook = load_workbook(filename=BytesIO(file_bytes), data_only=True)
    sheet = workbook.active

    imported = 0
    updated = 0
    errors: list[dict] = []

    site_map = {
        _normalize_text(site.nom): site
        for site in simulation.sites_travail.all()
    }
    vehicule_map = {
        _normalize_text(vehicule.immatriculation): vehicule
        for vehicule in simulation.vehicules.all()
    }

    for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if not row or all(value is None or str(value).strip() == "" for value in row[:3]):
            continue

        if index == 1 and isinstance(row[0], str) and _normalize_text(row[0]) == "date":
            continue

        try:
            raw_date, raw_site_name, raw_immatriculation = row[:3]
            jour_date = _parse_excel_date(raw_date)
            site_name = _normalize_text(str(raw_site_name or ""))
            immatriculation = _normalize_text(str(raw_immatriculation or ""))

            if not site_name or not immatriculation:
                raise ImportFraisKilometriquesError("Le nom du site et l immatriculation sont obligatoires.")

            site = site_map.get(site_name)
            if site is None:
                raise ImportFraisKilometriquesError(f"Site inconnu: {raw_site_name}")

            vehicule = vehicule_map.get(immatriculation)
            if vehicule is None:
                raise ImportFraisKilometriquesError(f"Vehicule inconnu: {raw_immatriculation}")

            if vehicule.date_achat and jour_date < vehicule.date_achat:
                raise ImportFraisKilometriquesError("Le vehicule ne peut pas etre utilise avant sa date d'achat.")
            if vehicule.date_vente and jour_date > vehicule.date_vente:
                raise ImportFraisKilometriquesError("Le vehicule ne peut pas etre utilise apres sa date de vente.")
            if not hasattr(simulation, "domicile"):
                raise ImportFraisKilometriquesError("Le domicile est obligatoire avant un import sur site.")

            try:
                distance = get_or_compute_distance(domicile=simulation.domicile, site=site)
            except OpenRouteServiceError as exc:
                raise ImportFraisKilometriquesError(str(exc)) from exc

            jour, created = JourTravaille.objects.update_or_create(
                simulation=simulation,
                date=jour_date,
                defaults={
                    "type_jour": JourTravaille.TypeJour.SITE,
                    "site_travail": site,
                    "vehicule": vehicule,
                    "distance_km": distance.distance_aller_retour_km,
                    "montant_eur": 0,
                    "commentaire": "Import Excel",
                },
            )

            if created:
                imported += 1
            else:
                updated += 1
        except ImportFraisKilometriquesError as exc:
            errors.append({"row": index, "detail": str(exc)})

    return {
        "imported": imported,
        "updated": updated,
        "errors": errors,
    }
